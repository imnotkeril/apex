"""Chart rendering, PPTX export and their endpoints."""

import uuid
import zipfile
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import HTTPException

from app.modules.content import router as content_router
from app.modules.content import router_ext
from app.modules.content.models import ContentDraft
from app.modules.content.schemas import DraftPatch
from app.modules.content.services import chart_render, pptx_export
from tests._fakes import FakeResult, FakeSession

ORG = uuid.uuid4()
BAR = {"type": "bar", "labels": ["Q1", "Q2"], "series": [{"name": "Rev", "values": [1, 2]}, {"name": "Cost", "values": [1, 1.5]}]}


@pytest.mark.parametrize("data", [
    BAR,
    {"type": "line", "labels": ["a", "b", "c"], "series": [{"name": "x", "values": [3, 1, 2]}]},
    {"type": "pie", "labels": ["a", "b"], "series": [{"values": [60, 40]}]},
])
def test_render_each_chart_type(data):
    png = chart_render.render_chart_png(data, {"title": "T", "colors": ["#ff0000", "#00ff00"]})
    assert png.startswith(b"\x89PNG")


@pytest.mark.parametrize("data,msg", [
    ("nope", "object"),
    ({"type": "radar", "labels": ["a"], "series": [{"values": [1]}]}, "Unsupported"),
    ({"type": "bar", "labels": [], "series": []}, "labels"),
    ({"type": "bar", "labels": ["a", "b"], "series": [{"values": [1]}]}, "one value per label"),
    ({"type": "bar", "labels": ["a"], "series": [{"values": ["x"]}]}, "numbers"),
    ({"type": "pie", "labels": ["a"], "series": [{"values": [1]}, {"values": [2]}]}, "exactly one"),
])
def test_invalid_chart_data(data, msg):
    with pytest.raises(chart_render.InvalidChartData, match=msg):
        chart_render.validate(data)


def test_split_body_respects_limit_and_keeps_text():
    body = "\n\n".join(["Sentence one. " * 20, "Short para.", "Word " * 300])
    chunks = pptx_export.split_body(body, limit=400)
    assert all(len(c) <= 400 for c in chunks) and len(chunks) > 3
    assert "Short para." in "\n".join(chunks)


def _slides(pptx: bytes) -> list[str]:
    with zipfile.ZipFile(BytesIO(pptx)) as z:
        return sorted(n for n in z.namelist() if n.startswith("ppt/slides/slide"))


def test_build_pptx_structure_and_unicode():
    from pptx import Presentation

    deck = pptx_export.build_pptx("Заголовок", "Первый абзац.\n\nВторой абзац.", ["RU"], chart_render.render_chart_png(BAR))
    prs = Presentation(BytesIO(deck))
    assert len(prs.slides) == 3  # cover + chart + 1 body
    assert prs.slides[0].shapes.title.text == "Заголовок"
    assert "Первый абзац." in prs.slides[2].placeholders[1].text_frame.text


def _draft(**kw) -> ContentDraft:
    return ContentDraft(id=uuid.uuid4(), org_id=ORG, topic="Topic", language="en", status="ready", styled_text="Body text",
                        translations={}, **kw)


@pytest.mark.asyncio
async def test_pptx_endpoint_with_and_without_chart():
    user = SimpleNamespace(org_id=ORG)
    d = _draft(chart_data=BAR)
    resp = await router_ext.draft_pptx(d.id, FakeSession([FakeResult(items=[d])]), user, lang=None)
    assert resp.media_type.endswith("presentationml.presentation") and len(_slides(resp.body)) == 3
    d2 = _draft()
    resp = await router_ext.draft_pptx(d2.id, FakeSession([FakeResult(items=[d2])]), user, lang=None)
    assert len(_slides(resp.body)) == 2


@pytest.mark.asyncio
async def test_chart_endpoint():
    user = SimpleNamespace(org_id=ORG)
    d = _draft(chart_data=BAR)
    resp = await router_ext.draft_chart(d.id, FakeSession([FakeResult(items=[d])]), user)
    assert resp.media_type == "image/png" and resp.body.startswith(b"\x89PNG")
    empty = _draft()
    with pytest.raises(HTTPException) as e:
        await router_ext.draft_chart(empty.id, FakeSession([FakeResult(items=[empty])]), user)
    assert e.value.status_code == 404
    bad = _draft(chart_data={"type": "bar", "labels": ["a"], "series": [{"values": [1, 2]}]})
    with pytest.raises(HTTPException) as e:
        await router_ext.draft_chart(bad.id, FakeSession([FakeResult(items=[bad])]), user)
    assert e.value.status_code == 422


@pytest.mark.asyncio
async def test_patch_draft_validates_chart_data():
    user = SimpleNamespace(org_id=ORG)
    d = _draft()
    out = await content_router.patch_draft(d.id, DraftPatch(chart_data=BAR), FakeSession([FakeResult(items=[d])]), user)
    assert out.chart_data == BAR
    with pytest.raises(HTTPException) as e:
        await content_router.patch_draft(d.id, DraftPatch(chart_data={"type": "bar"}), FakeSession([FakeResult(items=[d])]), user)
    assert e.value.status_code == 422
